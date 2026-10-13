#!/usr/bin/env python3
"""生成端到端校验用的「真实」样张。

所有文件均由与引擎无关的第三方库生成（python-docx / openpyxl /
python-pptx / fpdf2），用于校验 rust_doc_engine 的解析、渲染、合并能力。

用法：python3 samples/gen_samples.py
"""
import os

HERE = os.path.dirname(os.path.abspath(__file__))


def gen_docx() -> None:
    import docx

    d = docx.Document()
    d.add_heading("周报（端到端样张）", level=0)
    p = d.add_paragraph("这是由 python-docx 生成的示例文档，用于校验 Rust 引擎的 DOCX 解析。")
    p.add_run("这一句是加粗的。").bold = True
    d.add_paragraph("下面是一张 2x3 表格：")
    t = d.add_table(rows=2, cols=3)
    for r in range(2):
        for c in range(3):
            t.cell(r, c).text = f"R{r}C{c}"
    d.save(os.path.join(HERE, "hello.docx"))


def gen_xlsx() -> None:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "销售"
    ws.append(["商品", "数量", "单价"])
    ws.append(["苹果", 30, 3.5])
    ws.append(["香蕉", 12, 4.2])
    ws["D2"] = "=SUM(B2:B3)"
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws2 = wb.create_sheet("备注")
    ws2["A1"] = "纯文本说明"
    wb.save(os.path.join(HERE, "报表.xlsx"))


def gen_pptx() -> None:
    from pptx import Presentation

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "端到端校验·第一页"
    s1.placeholders[1].text = "要点一\n要点二\n要点三"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "第二页标题"
    s2.placeholders[1].text = "详情 甲\n详情 乙 含中文"
    prs.save(os.path.join(HERE, "演示.pptx"))


def gen_pdf() -> None:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_title("Sample PDF title")
    pdf.set_author("e2e-author")
    pdf.add_page()
    pdf.set_font("helvetica", "B", 22)
    pdf.cell(0, 16, "End-to-End Sample", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("helvetica", "", 12)
    pdf.multi_cell(
        0, 8,
        "Hello from a real-generated PDF. Line two of body text for verification "
        "of text extraction across the page.",
    )
    pdf.add_page()
    pdf.cell(0, 16, "Page Two", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("helvetica", "", 12)
    pdf.multi_cell(0, 8, "Second page content 42.7 plus numbers 12345.")
    pdf.output(os.path.join(HERE, "sample.pdf"))


if __name__ == "__main__":
    for f in (gen_docx, gen_xlsx, gen_pptx, gen_pdf):
        f()
    print("generated:", ", ".join(sorted(os.listdir(HERE))))